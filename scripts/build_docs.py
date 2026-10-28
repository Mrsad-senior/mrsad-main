"""Build the IS499 Word deliverables from Markdown sources.

Each document is described by a YAML manifest in ``docs/report/`` (template, cover fields,
front matter and the ordered list of Markdown files). The Markdown dialect is deliberately
small:

* ``#``..``####`` headings (``#`` starts a new page), paragraphs, ``-`` bullets (two-space nesting),
  ``1.`` numbered items, ``>`` notes, fenced code, pipe tables, ``<!-- pagebreak -->``.
* Figures: ``![Caption](path.png){#fig:label width=5.5}``; tables: a ``Table: Caption {#tab:label}``
  line directly above the table. ``@fig:label`` / ``@tab:label`` become "Figure 3.2" / "Table 4.1".
* ``::: rtl`` ... ``:::`` wraps right-to-left Arabic paragraphs.
* Inline ``**bold**``, ``*italic*``, `` `code` `` and ``[text](url)``.

Usage: python scripts/build_docs.py docs/report/final_report.yaml [...more manifests]
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
NAVY = RGBColor(0x17, 0x36, 0x5D)
BLUE = RGBColor(0x36, 0x5F, 0x91)
GREY = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "1F3864"
ZEBRA_FILL = "EEF3F8"
BODY_FONT = "Times New Roman"
ARABIC_FONT = "Arial"
MONO_FONT = "Consolas"

INLINE_RE = re.compile(
    r"(\[@[\w-]+(?:\s*[;,]\s*@[\w-]+)*\]|\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|@(?:fig|tab):[\w-]+)"
)
CITE_RE = re.compile(r"\[(@[\w-]+(?:\s*[;,]\s*@[\w-]+)*)\]")
FIG_RE = re.compile(r"^!\[(?P<cap>[^\]]*)\]\((?P<path>[^)]+)\)(?:\{(?P<attrs>[^}]*)\})?\s*$")
TABCAP_RE = re.compile(r"^Table:\s*(?P<cap>.*?)\s*(?:\{(?P<attrs>[^{}]*)\})?\s*$")


def table_attrs(attrs: str | None) -> tuple[str | None, list[float] | None]:
    """Parse ``{#tab:label widths=1,3,2}`` (either part optional) from a table caption line."""
    if not attrs:
        return None, None
    label = re.search(r"#(tab:[\w-]+)", attrs)
    widths = re.search(r"widths=([\d.,\s]+)", attrs)
    return (label.group(1) if label else None,
            [float(x) for x in widths.group(1).replace(" ", "").split(",") if x] if widths else None)


# --------------------------------------------------------------------------- low-level XML helpers


def add_field(paragraph, instr: str, placeholder: str = "") -> None:
    """Insert a complex field (TOC, PAGE, …); Word computes its result when fields are updated."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._r.append(begin)
    run = paragraph.add_run()
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instr} "
    run._r.append(text)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=90, right=90) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_rtl(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    p_pr.insert(0, bidi)
    for run in paragraph.runs:
        r_pr = run._r.get_or_add_rPr()
        r_pr.append(OxmlElement("w:rtl"))
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, fonts)
        fonts.set(qn("w:cs"), ARABIC_FONT)
        size = OxmlElement("w:szCs")
        size.set(qn("w:val"), "26")
        r_pr.append(size)


def paragraph_shading(paragraph, fill: str, border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        bdr.append(left)
        p_pr.append(bdr)


def page_number_format(section, fmt: str, start: int | None = 1) -> None:
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sect_pr.append(pg)


# --------------------------------------------------------------------------- numbering pass


@dataclass
class Labels:
    figures: dict[str, str] = field(default_factory=dict)
    tables: dict[str, str] = field(default_factory=dict)
    citations: dict[str, int] = field(default_factory=dict)


def chapter_key(heading: str, counters: dict) -> str | None:
    m = re.match(r"(?:Chapter|الفصل)\s+(\d+)", heading)
    if m:
        return m.group(1)
    m = re.match(r"Appendix\s+([A-Z])", heading)
    if m:
        return m.group(1)
    return None


def caption_number(chapter: str, n: int) -> str:
    """"3.2" inside chapters/appendices; plain "2" in documents (or front matter) without chapters."""
    return f"{chapter}.{n}" if chapter != "0" else str(n)


def collect_labels(markdown_texts: list[str]) -> Labels:
    labels = Labels()
    chapter = "0"
    fig_n = tab_n = 0
    in_code = False
    for text in markdown_texts:
        for line in text.splitlines():
            if line.startswith("```"):
                in_code = not in_code
                continue
            if in_code:
                continue
            for cm in CITE_RE.finditer(line):
                for key in re.findall(r"@([\w-]+)", cm.group(1)):
                    labels.citations.setdefault(key, len(labels.citations) + 1)
            if line.startswith("# "):
                key = chapter_key(line[2:].strip(), {})
                if key:
                    chapter, fig_n, tab_n = key, 0, 0
                continue
            m = FIG_RE.match(line.strip())
            if m:
                fig_n += 1
                attrs = m.group("attrs") or ""
                lm = re.search(r"#(fig:[\w-]+)", attrs)
                if lm:
                    labels.figures[lm.group(1)] = caption_number(chapter, fig_n)
                continue
            m = TABCAP_RE.match(line.strip())
            if m:
                tab_n += 1
                label, _ = table_attrs(m.group("attrs"))
                if label:
                    labels.tables[label] = caption_number(chapter, tab_n)
    return labels


VALUE_RE = re.compile(r"\{\{(?P<file>[^#{}]+)#(?P<path>[^|{}]+)(?:\|(?P<fmt>[^{}]+))?\}\}")
PENDING = "[pending]"


def lookup_value(file: Path, dotted: str, fmt: str | None, cache: dict) -> str:
    """Read a number/string from a JSON result file: ``{{reports/x.json#a.b.0.c|.2f}}``."""
    import json

    if file not in cache:
        cache[file] = json.loads(file.read_text(encoding="utf-8")) if file.exists() else None
    node = cache[file]
    for part in dotted.split("."):
        if node is None:
            return PENDING
        if isinstance(node, list):
            try:
                node = node[int(part)]
            except (ValueError, IndexError):
                return PENDING
        elif isinstance(node, dict):
            node = node.get(part)
        else:
            return PENDING
    if node is None:
        return PENDING
    if fmt:
        try:
            return format(node, fmt)
        except (TypeError, ValueError):
            return str(node)
    return str(node)


INCLUDE_RE = re.compile(r"^<<include\s+(?P<path>[^>\s]+)(?P<nohead>\s+nohead)?\s*>>\s*$", re.M)


def expand_includes(text: str, resolve, depth: int = 0) -> str:
    """Inline ``<<include path.md>>`` directives (recursively, max depth 5).

    ``<<include path.md nohead>>`` drops the included file's first ``# `` heading, so a chapter can be
    reused inside a standalone document that has its own title.
    """
    if depth > 5:
        raise RecursionError("include depth exceeded")

    def sub(m: re.Match) -> str:
        body = resolve(m.group("path").strip()).read_text(encoding="utf-8")
        if m.group("nohead"):
            body = re.sub(r"^# .*\n", "", body, count=1, flags=re.M)
        return expand_includes(body, resolve, depth + 1)

    return INCLUDE_RE.sub(sub, text)


# --------------------------------------------------------------------------- builder


class DocBuilder:
    def __init__(self, manifest_path: Path):
        self.manifest_path = manifest_path
        self.m = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
        self.base = manifest_path.parent
        self.doc = Document(str(self._path(self.m["template"])))
        self.chapter = "0"
        self.fig_n = 0
        self.tab_n = 0
        self.labels = Labels()
        self.missing_refs: list[str] = []
        self.pending_values: list[str] = []
        bib_path = self.m.get("bibliography")
        self.bib: dict[str, str] = {}
        if bib_path:
            for entry in yaml.safe_load(self._path(bib_path).read_text(encoding="utf-8")) or []:
                self.bib[entry["key"]] = entry["ieee"]
        self._clear_body()
        self._styles()

    def _read(self, p: str) -> str:
        text = expand_includes(self._path(p).read_text(encoding="utf-8"), self._path)
        cache: dict = {}

        def sub(m: re.Match) -> str:
            value = lookup_value(ROOT / m.group("file").strip(), m.group("path").strip(), m.group("fmt"), cache)
            if value == PENDING:
                self.pending_values.append(m.group(0))
            return value

        return VALUE_RE.sub(sub, text)

    def _path(self, p: str) -> Path:
        path = Path(p)
        if path.is_absolute():
            return path
        cand = (self.base / path).resolve()
        return cand if cand.exists() else (ROOT / path).resolve()

    # ---------------------------------------------------------------- setup

    def _clear_body(self) -> None:
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)

    def _styles(self) -> None:
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        normal.element.rPr.rFonts.set(qn("w:cs"), ARABIC_FONT)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13.5, BLUE), ("Heading 3", 12, BLUE),
                                  ("Heading 4", 12, GREY)):
            s = st[name]
            s.font.name = BODY_FONT
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.color.rgb = color
            rpr = s.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.insert(0, fonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(attr), BODY_FONT)
            for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if fonts.get(qn(theme_attr)) is not None:
                    del fonts.attrib[qn(theme_attr)]
            s.paragraph_format.space_before = Pt(18 if name == "Heading 1" else 12)
            s.paragraph_format.space_after = Pt(8 if name == "Heading 1" else 4)
            s.paragraph_format.keep_with_next = True
        for name in ("Figure Caption", "Table Caption"):
            if name not in [s.name for s in st]:
                cap = st.add_style(name, 1)
                cap.base_style = st["Caption"]
                cap.font.name = BODY_FONT
                cap.font.size = Pt(10)
                cap.font.bold = True
                cap.font.italic = False
                cap.font.color.rgb = NAVY
                cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_before = Pt(4)
                cap.paragraph_format.space_after = Pt(10)
        if "Code Block" not in [s.name for s in st]:
            code = st.add_style("Code Block", 1)
            code.base_style = st["Normal"]
            code.font.name = MONO_FONT
            code.font.size = Pt(8.5)
            code.paragraph_format.space_after = Pt(8)
            code.paragraph_format.line_spacing = 1.0
        if "Table Text" not in [s.name for s in st]:
            tt = st.add_style("Table Text", 1)
            tt.base_style = st["Normal"]
            tt.font.size = Pt(9.5)
            tt.paragraph_format.space_after = Pt(0)
            tt.paragraph_format.line_spacing = 1.0

    # ---------------------------------------------------------------- inline

    def _inline(self, paragraph, text: str, *, bold: bool = False, size: float | None = None,
                color: RGBColor | None = None) -> None:
        for part in INLINE_RE.split(text):
            if not part:
                continue
            run_bold, run_italic, mono = bold, False, False
            if part.startswith("[@"):
                nums = []
                for key in re.findall(r"@([\w-]+)", part):
                    n = self.labels.citations.get(key)
                    if n is None or key not in self.bib:
                        self.missing_refs.append("cite:" + key)
                    nums.append(f"[{n if n is not None else '?'}]")
                part = ", ".join(nums)
            elif part.startswith("**") and part.endswith("**") and len(part) > 4:
                part, run_bold = part[2:-2], True
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                part, mono = part[1:-1], True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                part, run_italic = part[1:-1], True
            elif part.startswith("[") and "](" in part:
                label, url = part[1:].split("](", 1)
                url = url[:-1]
                part = label if label == url or url.startswith("#") else f"{label} ({url})"
            elif part.startswith("@fig:") or part.startswith("@tab:"):
                key = part[1:]
                num = (self.labels.figures if key.startswith("fig:") else self.labels.tables).get(key)
                if num is None:
                    self.missing_refs.append(key)
                    num = "??"
                part = ("Figure " if key.startswith("fig:") else "Table ") + num
            run = paragraph.add_run(part)
            run.bold = run_bold or None
            run.italic = run_italic or None
            if mono:
                run.font.name = MONO_FONT
                run.font.size = Pt(9.5 if size is None else max(size - 1, 7))
                run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            elif size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color

    # ---------------------------------------------------------------- blocks

    def heading(self, text: str, level: int) -> None:
        if level == 1:
            key = chapter_key(text, {})
            if key:
                self.chapter, self.fig_n, self.tab_n = key, 0, 0
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        if level == 1 and self._has_content():
            p.paragraph_format.page_break_before = True
        self._inline(p, text)
        if re.search(r"[؀-ۿ]", text) and not re.search(r"[A-Za-z]", text):
            set_rtl(p)

    def _has_content(self) -> bool:
        return len(self.doc.element.body) > 1

    def paragraph(self, text: str, *, rtl: bool = False, align=None) -> None:
        p = self.doc.add_paragraph()
        self._inline(p, text)
        p.paragraph_format.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
        if rtl:
            set_rtl(p)

    def bullet(self, text: str, level: int) -> None:
        style = "List Bullet" if level == 0 else "List Bullet 2" if level == 1 else "List Bullet 3"
        p = self.doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        self._inline(p, text)

    def numbered(self, number: str, text: str, level: int) -> None:
        p = self.doc.add_paragraph()
        indent = 0.35 + 0.3 * level
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(indent))
        self._inline(p, f"{number}\t{text}")

    def note(self, text: str) -> None:
        p = self.doc.add_paragraph()
        paragraph_shading(p, "EEF6FB", border="2E75B6")
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.right_indent = Inches(0.1)
        self._inline(p, text, size=11)

    def code(self, lines: list[str]) -> None:
        p = self.doc.add_paragraph(style="Code Block")
        paragraph_shading(p, "F3F4F6", border="9CA3AF")
        for i, line in enumerate(lines):
            run = p.add_run(line.replace("\t", "    ") or " ")
            run.font.name = MONO_FONT
            if i < len(lines) - 1:
                run.add_break()

    def figure(self, caption: str, path: str, attrs: str) -> None:
        self.fig_n += 1
        width = 6.0
        wm = re.search(r"width=([\d.]+)", attrs or "")
        if wm:
            width = min(float(wm.group(1)), 6.0)
        img = self._path(path)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        if img.exists():
            p.add_run().add_picture(str(img), width=Inches(width))
        else:
            r = p.add_run(f"[missing figure: {path}]")
            r.font.color.rgb = RGBColor(0xC0, 0, 0)
            print(f"  WARNING missing figure {img}")
        cap = self.doc.add_paragraph(style="Figure Caption")
        self._inline(cap, f"Figure {caption_number(self.chapter, self.fig_n)}: {caption}")

    def table(self, rows: list[list[str]], caption: str | None, widths: list[float] | None = None) -> None:
        if caption is not None:
            self.tab_n += 1
            cap = self.doc.add_paragraph(style="Table Caption")
            cap.paragraph_format.keep_with_next = True
            self._inline(cap, f"Table {caption_number(self.chapter, self.tab_n)}: {caption}")
        header, body = rows[0], rows[1:]
        ncols = len(header)
        table = self.doc.add_table(rows=1 + len(body), cols=ncols)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_margins(table)
        for c, text in enumerate(header):
            cell = table.rows[0].cells[c]
            shade(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            p.style = self.doc.styles["Table Text"]
            self._inline(p, text.strip(), bold=True, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF))
        repeat_header(table.rows[0])
        for r, row in enumerate(body, start=1):
            for c in range(ncols):
                text = row[c].strip() if c < len(row) else ""
                cell = table.rows[r].cells[c]
                if r % 2 == 0:
                    shade(cell, ZEBRA_FILL)
                lines = text.split("<br>")
                p = cell.paragraphs[0]
                p.style = self.doc.styles["Table Text"]
                self._inline(p, lines[0], size=9.5)
                for extra in lines[1:]:
                    q = cell.add_paragraph(style="Table Text")
                    self._inline(q, extra, size=9.5)
                if re.search(r"[؀-ۿ]", text) and not re.search(r"[A-Za-z]{3}", text):
                    for para in cell.paragraphs:
                        set_rtl(para)
        if widths:
            total = sum(widths)
            for c, w in enumerate(widths):
                for row in table.rows:
                    row.cells[c].width = Inches(6.0 * w / total)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def references(self) -> None:
        for key, n in sorted(self.labels.citations.items(), key=lambda kv: kv[1]):
            text = self.bib.get(key, f"MISSING REFERENCE {key}")
            text = re.sub(r"^\[\d+\]\s*", "", text.strip())
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.first_line_indent = Inches(-0.45)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
            self._inline(p, f"[{n}]\t{text}", size=10.5)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- markdown

    def render_markdown(self, text: str) -> None:
        lines = text.splitlines()
        i = 0
        para: list[str] = []
        rtl = False
        pending_caption: tuple[str, list[float] | None] | None = None

        def flush() -> None:
            nonlocal para
            if para:
                self.paragraph(" ".join(s.strip() for s in para), rtl=rtl)
                para = []

        while i < len(lines):
            line = lines[i]
            stripped = line.strip()
            if stripped.startswith("<!--"):
                flush()
                if "pagebreak" in stripped:
                    self.page_break()
                while "-->" not in lines[i]:
                    i += 1
                i += 1
                continue
            if stripped == "<<references>>":
                flush()
                self.references()
                i += 1
                continue
            if stripped == "::: rtl":
                flush()
                rtl = True
                i += 1
                continue
            if stripped == ":::":
                flush()
                rtl = False
                i += 1
                continue
            if stripped.startswith("```"):
                flush()
                block = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    block.append(lines[i])
                    i += 1
                self.code(block)
                i += 1
                continue
            hm = re.match(r"^(#{1,4})\s+(.*)$", line)
            if hm:
                flush()
                self.heading(hm.group(2).strip(), len(hm.group(1)))
                i += 1
                continue
            fm = FIG_RE.match(stripped)
            if fm:
                flush()
                self.figure(fm.group("cap"), fm.group("path"), fm.group("attrs") or "")
                i += 1
                continue
            tm = TABCAP_RE.match(stripped)
            if tm and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
                flush()
                _label, widths = table_attrs(tm.group("attrs"))
                pending_caption = (tm.group("cap"), widths)
                i += 1
                continue
            if stripped.startswith("|"):
                flush()
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c for c in re.split(r"(?<!\\)\|", lines[i].strip())[1:-1]]
                    cells = [c.replace("\\|", "|") for c in cells]
                    if not all(re.fullmatch(r"\s*:?-{2,}:?\s*", c) for c in cells):
                        rows.append(cells)
                    i += 1
                caption, widths = pending_caption if pending_caption else (None, None)
                self.table(rows, caption, widths)
                pending_caption = None
                continue
            bm = re.match(r"^(\s*)[-*]\s+(.*)$", line)
            if bm and not stripped.startswith("**"):
                flush()
                self.bullet(bm.group(2), len(bm.group(1)) // 2)
                i += 1
                continue
            nm = re.match(r"^(\s*)(\d+[.)])\s+(.*)$", line)
            if nm:
                flush()
                self.numbered(nm.group(2).rstrip(")").rstrip(".") + ".", nm.group(3), len(nm.group(1)) // 3)
                i += 1
                continue
            if stripped.startswith(">"):
                flush()
                quote = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    quote.append(lines[i].strip()[1:].strip())
                    i += 1
                self.note(" ".join(quote))
                continue
            if not stripped:
                flush()
                i += 1
                continue
            para.append(line)
            i += 1
        flush()

    # ---------------------------------------------------------------- document parts

    def cover(self) -> None:
        c = self.m["cover"]
        logo = self._path(c.get("logo", "docs/report/assets/psu_logo.png"))
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if logo.exists():
            p.add_run().add_picture(str(logo), width=Inches(2.6))
        for line, size, bold in ((c.get("university", "Prince Sultan University"), 15, True),
                                 (c.get("college", "College of Computer & Information Sciences"), 12.5, False),
                                 (c.get("department", "Information Systems Department"), 12.5, False),
                                 (c.get("course", "IS499 – Senior Project"), 12.5, False)):
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(2)
            r = q.add_run(line)
            r.font.size = Pt(size)
            r.bold = bold
            r.font.color.rgb = NAVY
        q = self.doc.add_paragraph()
        q.paragraph_format.space_before = Pt(30)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = q.add_run(c["doc_type"])
        r.font.size = Pt(13)
        r.font.color.rgb = GREY
        r.bold = True
        q = self.doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(4)
        r = q.add_run(c["title"])
        r.font.size = Pt(22)
        r.bold = True
        r.font.color.rgb = NAVY
        if c.get("subtitle"):
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = q.add_run(c["subtitle"])
            r.font.size = Pt(13.5)
            r.italic = True
            r.font.color.rgb = BLUE
        q = self.doc.add_paragraph()
        q.paragraph_format.space_before = Pt(24)
        rows = [["Team member", "Student ID", "Department"]] + [[m["name"], str(m["id"]), m["dept"]] for m in c["team"]]
        self.table(rows, None, [3, 1.6, 1.4])
        info = c.get("info", [])
        t = self.doc.add_table(rows=len(info), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_i, (k, v) in enumerate(info):
            a, b = t.rows[r_i].cells
            a.width, b.width = Inches(2.2), Inches(3.8)
            pa = a.paragraphs[0]
            ra = pa.add_run(k)
            ra.bold = True
            ra.font.size = Pt(11)
            ra.font.color.rgb = NAVY
            rb = b.paragraphs[0].add_run(v)
            rb.font.size = Pt(11)

    def front_list(self, title: str, instr: str, placeholder: str, *, in_toc: bool = True) -> None:
        p = self.doc.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = True
        if not in_toc:
            # Keep the Heading 1 look but exclude the paragraph from the TOC field (\u uses outline levels).
            p_pr = p._p.get_or_add_pPr()
            lvl = OxmlElement("w:outlineLvl")
            lvl.set(qn("w:val"), "9")
            p_pr.append(lvl)
        self._inline(p, title)
        q = self.doc.add_paragraph()
        add_field(q, instr, placeholder)

    def new_section(self, fmt: str, start: int | None, header_text: str | None) -> None:
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        page_number_format(sec, fmt, start)
        sec.footer.is_linked_to_previous = False
        sec.header.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, "PAGE", "1")
        hp = sec.header.paragraphs[0]
        hp.text = ""
        if header_text:
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = hp.add_run(header_text)
            r.font.size = Pt(9)
            r.italic = True
            r.font.color.rgb = GREY

    def build(self, out: Path) -> None:
        texts = []
        for item in self.m.get("front", []) + self.m["body"]:
            if isinstance(item, str):
                texts.append(self._read(item))
        self.labels = collect_labels(texts)
        first = self.doc.sections[0]
        first.different_first_page_header_footer = False
        if self.m.get("cover"):
            self.cover()
        header_text = self.m.get("running_header")
        self.new_section("lowerRoman", 1, None)
        for item in self.m.get("front", []):
            if isinstance(item, dict) and "list" in item:
                kind = item["list"]
                if kind == "toc":
                    self.front_list("Table of Contents", 'TOC \\o "1-3" \\h \\z \\u',
                                    "Right-click and choose Update Field to build the table of contents.", in_toc=False)
                elif kind == "figures":
                    self.front_list("List of Figures", 'TOC \\h \\z \\t "Figure Caption,1"',
                                    "Right-click and choose Update Field to list the figures.")
                elif kind == "tables":
                    self.front_list("List of Tables", 'TOC \\h \\z \\t "Table Caption,1"',
                                    "Right-click and choose Update Field to list the tables.")
                continue
            self.render_markdown(self._read(item))
        self.new_section("decimal", 1, header_text)
        for item in self.m["body"]:
            self.render_markdown(self._read(item))
        self._fix_first_heading_breaks()
        if self.m.get("update_fields_on_open", False):
            upd = OxmlElement("w:updateFields")
            upd.set(qn("w:val"), "true")
            self.doc.settings.element.append(upd)
        props = self.doc.core_properties
        props.title = self.m["cover"]["title"] if self.m.get("cover") else out.stem
        props.author = "Mirsad team – Prince Sultan University"
        props.subject = self.m.get("subject", "IS499 Senior Project")
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        if self.missing_refs:
            print("  WARNING unresolved references:", sorted(set(self.missing_refs)))
        if self.pending_values:
            print(f"  NOTE {len(set(self.pending_values))} result values pending (result files not generated yet)")
        print(f"wrote {out}  (figures/tables labelled: {len(self.labels.figures)}/{len(self.labels.tables)})")

    def _fix_first_heading_breaks(self) -> None:
        """Remove page-break-before on headings that directly follow a section break."""
        body = self.doc.element.body
        prev_was_sectbreak = True
        for el in body:
            if el.tag != qn("w:p"):
                prev_was_sectbreak = False
                continue
            p_pr = el.find(qn("w:pPr"))
            has_sect = p_pr is not None and p_pr.find(qn("w:sectPr")) is not None
            if prev_was_sectbreak and p_pr is not None:
                pbb = p_pr.find(qn("w:pageBreakBefore"))
                if pbb is not None:
                    p_pr.remove(pbb)
            text = "".join(t.text or "" for t in el.iter(qn("w:t")))
            prev_was_sectbreak = has_sect or (prev_was_sectbreak and not text.strip() and p_pr is None)


def main(argv: list[str]) -> None:
    for manifest in argv:
        mpath = Path(manifest).resolve()
        builder = DocBuilder(mpath)
        out = ROOT / builder.m["output"]
        builder.build(out)


if __name__ == "__main__":
    main(sys.argv[1:] or [str(ROOT / "docs/report/final_report.yaml")])
