"""Render a PowerPoint deck to one PNG per slide with macOS Quick Look (no PowerPoint automation needed).

Quick Look's thumbnailer renders only the first slide, so each slide is saved as its own single-slide copy
first. Quick Look also hangs on decks whose slides have speaker notes created by python-pptx, so the
preview copies drop the notes (the real deck keeps them). Usage: python scripts/preview_slides.py deck.pptx out_dir [--size 1600]
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation


def single_slide_copy(src: Path, index: int, dest: Path) -> None:
    prs = Presentation(str(src))
    ids = prs.slides._sldIdLst  # noqa: SLF001 - python-pptx has no public slide-removal API
    for i, sld in reversed(list(enumerate(list(ids)))):
        if i != index:
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)
    slide = prs.slides[0]
    for rel in list(slide.part.rels.values()):
        if rel.reltype.endswith("/notesSlide"):
            slide.part.drop_rel(rel.rId)
    prs.save(str(dest))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("out")
    ap.add_argument("--size", type=int, default=1600)
    args = ap.parse_args()
    deck, out = Path(args.deck), Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    count = len(Presentation(str(deck)).slides)
    with tempfile.TemporaryDirectory() as tmp:
        for i in range(count):
            single = Path(tmp) / f"slide-{i + 1:02d}.pptx"
            single_slide_copy(deck, i, single)
            subprocess.run(["qlmanage", "-t", "-s", str(args.size), "-o", tmp, str(single)],  # noqa: S603,S607
                           check=False, capture_output=True, timeout=90)
            png = Path(tmp) / f"{single.name}.png"
            if png.exists():
                shutil.move(str(png), out / f"slide-{i + 1:02d}.png")
                print("rendered", i + 1)
            else:
                print("no preview for slide", i + 1)


if __name__ == "__main__":
    main()
